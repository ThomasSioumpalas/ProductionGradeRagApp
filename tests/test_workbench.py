import asyncio
import json
import os
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient

from financial_workbench.api import create_app
from financial_workbench.analysis import (check_reported_measures, enrich_financial_evidence,
                                          kpi_coverage, merge_evidence)
from financial_workbench.documents import page_windows, read_pdf, statement_plan
from financial_workbench.engine import (
    CATALOG,
    METRICS,
    extract,
    mapping_tasks,
    parse_number,
    provisional_decisions,
    reconcile,
    validate_fact,
)
from financial_workbench.models import ExtractedFact, Settings
from financial_workbench.store import Store
from financial_workbench.workbook import export_workbook
from financial_workbench.forecast import operating_scenarios
from financial_workbench.investigation import investigate_gaps


@pytest.fixture
def settings():
    return Settings(company="Example SA", latest_year=2025)


@pytest.fixture
def page():
    return {
        "id": 0,
        "file": "annual.pdf",
        "sha256": "a" * 64,
        "page": 1,
        "text": "Example SA. Consolidated. Annual 2025. EUR thousands. Statement of Profit or Loss.\nRevenue 1,234.5\nCost of sales (234.5)",
        "tables": "",
    }


def fact(**overrides):
    return ExtractedFact(
        **(
            {
                "metric_id": "income_statement_8",
                "year": 2025,
                "scope": "consolidated",
                "currency": "EUR",
                "raw_value": "1,234.5",
                "decimal_separator": ".",
                "scale": 1000,
                "source_file": "annual.pdf",
                "page": 1,
                "quote": "Revenue 1,234.5",
                "context_quote": "Annual 2025. EUR thousands.",
            }
            | overrides
        )
    )


@pytest.mark.parametrize(
    "raw,separator,expected",
    [
        ("1,234.56", ".", "1234.56"),
        ("1.234,56", ",", "1234.56"),
        ("(1.234,56)", ",", "-1234.56"),
        ("0", ".", "0"),
        ("−12,5", ",", "-12.5"),
        ("1 234,5", ",", "1234.5"),
    ],
)
def test_numeric_locales(raw, separator, expected):
    assert parse_number(raw, separator) == Decimal(expected)


@pytest.mark.parametrize(
    "raw", ["11.625.5", "331\\", "NaN", "inf", "=2+2", "1,23", "(-2)"]
)
def test_bad_numbers(raw):
    with pytest.raises(ValueError):
        parse_number(raw, ".")


def test_grounding_and_scale(settings, page):
    settings.company = "Any Excel display label"
    c = validate_fact(fact(), page, settings)
    assert c["value"] == "1.2345"
    assert c["file"] == "annual.pdf"
    assert (
        validate_fact(
            fact(
                metric_id="income_statement_9",
                raw_value="(234.5)",
                quote="Cost of sales (234.5)",
            ),
            page,
            settings,
        )["value"]
        == "0.2345"
    )


def test_tax_expense_sign_matches_workbook_convention(settings, page):
    source = {**page, "text": page["text"] + " Income taxes (176,435)"}
    c = validate_fact(fact(metric_id="income_statement_26", raw_value="(176,435)",
                           quote="Income taxes (176,435)"), source, settings)
    assert c["value"] == "176.435"


@pytest.mark.parametrize(
    "change",
    [
        {"quote": "Revenue 5", "raw_value": "5"},
        {"scope": "standalone"},
        {"currency": "USD"},
        {"year": 2018},
        {"page": 2},
        {"source_file": "another.pdf"},
        {"metric_id": "market_inputs_18"},
        {"raw_value": "234.5"},
    ],
)
def test_reject_unsupported_claim(settings, page, change):
    with pytest.raises(ValueError):
        validate_fact(fact(**change), page, settings)


def test_conflicts_are_kept_for_review(settings):
    def entry(n, raw):
        quote = f"Revenue {raw} {raw} 100 100"
        page = {"id": n, "file": "annual.pdf", "sha256": "a" * 64, "page": n,
                "text": f"GROUP COMPANY 31/12/2025 31/12/2024 {quote}", "tables": quote}
        row = {"row_id": f"{n}:full:0", "page": n, "file": "annual.pdf", "panel": "full",
               "statement": "income", "section": "", "label": "Revenue", "quote": quote,
               "currency": "EUR", "scale": 1000,
               "values": [{"year": 2025, "scope": "consolidated", "header": "31/12/2025",
                           "raw_value": raw, "x": 300}]}
        return page, row

    pairs = [entry(1, "1,234.5"), entry(2, "2,234.5")]
    panels = [{"rows": [row]} for _, row in pairs]
    progress = []
    candidates, rejected = asyncio.run(extract(
        [p for p, _ in pairs], settings, lambda a, b: progress.append((a, b)),
        plan=(panels, [[row] for _, row in pairs]),
    ))
    assert len(candidates) == 2 and all(c["status"] == "conflict" for c in candidates)
    assert not rejected and progress[-1] == (0, 0)


def test_draft_chooses_only_uncontested_figures(settings, page):
    a = validate_fact(fact(), page, settings)
    repeated = {**a, "id": "other-citation", "value": "1.23450"}
    conflict = {**a, "id": "conflicting-source", "value": "2.0"}
    other = {**a, "id": "other-metric", "metric_id": "income_statement_9", "value": "0.2"}
    assert [d["id"] for d in provisional_decisions([a, repeated, conflict, other])] == ["other-metric"]
    assert len(provisional_decisions([a, repeated])) == 1


def test_identically_named_reports_keep_distinct_document_citations(settings):
    pages, panels = [], []
    for digest, amount in (("a" * 64, "1,000"), ("b" * 64, "2,000")):
        quote = f"Revenue {amount}"
        pages.append({"sha256": digest, "file": "annual.pdf", "page": 1,
                      "text": "2025 " + quote, "tables": quote})
        panels.append({"rows": [{"row_id": "1:full:0", "document_id": digest,
                                   "file": "annual.pdf", "page": 1, "panel": "full",
                                   "statement": "income", "section": "", "label": "Revenue",
                                   "quote": quote, "currency": "EUR", "scale": 1000,
                                   "values": [{"year": 2025, "scope": "consolidated",
                                               "header": "2025", "raw_value": amount, "x": 330}]}]})
    candidates, rejected = asyncio.run(extract(
        pages, settings, lambda *_: None, plan=(panels, [p["rows"] for p in panels])))
    assert not rejected
    assert {(c["document_id"], c["value"]) for c in candidates} == {
        ("a" * 64, "1.000"), ("b" * 64, "2.000")}


def test_incomplete_label_mapping_is_split_and_retried(settings, page):
    from financial_workbench.llm import IncompleteOutputError
    calls, progress = [], []
    rows = [
        {"row_id": f"1:full:{i}", "page": 1, "file": "annual.pdf", "panel": "full",
         "statement": "income", "section": "", "label": label,
         "quote": f"{label} 1,234.5", "currency": "EUR", "scale": 1000,
         "values": [{"year": 2025, "scope": "consolidated", "header": "2025",
                     "raw_value": "1,234.5", "x": 200}]}
        for i, label in enumerate(["Revenue", "Other annual income", "Miscellaneous operating gain"])
    ]
    page = {**page, "text": page["text"] + " 2025 Revenue 1,234.5"}
    plan = ([{"rows": rows}], [rows])

    async def fake(_messages, structured=False, **_kwargs):
        calls.append(_messages)
        body = json.loads(_messages[1]["content"])
        if len(body["rows"]) > 1:
            raise IncompleteOutputError("length", "test output limit")
        return json.dumps({"mappings": []})

    candidates, rejected = asyncio.run(
        extract([page], settings, lambda a, b: progress.append((a, b)), complete=fake, plan=plan)
    )
    assert len(candidates) == 1 and rejected == []
    assert len(calls) == 3
    assert progress[-1] == (2, 2)


def test_pdf_and_page_chunks(tmp_path):
    p = tmp_path / "report.pdf"
    d = pymupdf.open()
    d.new_page().insert_text(
        (40, 60),
        "Example SA annual consolidated report 2025. EUR thousands. Statement of Profit or Loss. Revenue 100.",
    )
    d.save(p)
    d.close()
    pages, _warnings = read_pdf(p, "report.pdf")
    assert pages[0]["page"] == 1 and "Revenue 100" in pages[0]["text"]
    windows = list(page_windows([{**pages[0], "text": "x" * 35000}]))
    assert len(windows) == 16 and all(p["page"] == 1 for p, _ in windows)
    blank = tmp_path / "scan.pdf"
    d = pymupdf.open()
    d.new_page()
    d.save(blank)
    d.close()
    with pytest.raises(ValueError, match="OCR"):
        read_pdf(blank, "scan.pdf")


@pytest.mark.parametrize("pdf_env,expected_revenue,expected_operating,expected_page", [
    ("WORKBENCH_UNILEVER_PDF", "50503", "9037", 131),
    ("WORKBENCH_ADIDAS_PDF", "24811", "2056", 4),
])
def test_official_cross_company_statements_and_revenue_scenarios(
        tmp_path, pdf_env, expected_revenue, expected_operating, expected_page):
    """Real 2025 reports: varying year headers, layouts and page positions."""
    if not os.environ.get(pdf_env):
        pytest.skip(f"Set {pdf_env} to an official downloaded annual report")
    path = Path(os.environ[pdf_env])
    pages, _ = read_pdf(path, path.name)
    panels, batches = statement_plan(pages)
    assert panels and batches

    async def mapping(messages, **_):
        return '{"mappings":[]}'

    values, rejected = asyncio.run(extract(pages, Settings(company="Unrelated spreadsheet label",
        latest_year=2025), lambda *_: None, complete=mapping, plan=(panels, batches)))
    assert not rejected
    def selected(metric):
        return [c for c in values if c["year"] == 2025 and c["metric_id"] == metric]
    assert [(c["value"], c["page"]) for c in selected("income_statement_8")] == [
        (expected_revenue, expected_page)]
    assert [(c["value"], c["page"]) for c in selected("income_statement_16")] == [
        (expected_operating, expected_page)]
    if pdf_env == "WORKBENCH_UNILEVER_PDF":
        assert len(pages) >= 280 and expected_page > len(pages) // 3
        result = operating_scenarios(values, Settings(company="Excel only", latest_year=2025))
        assert result["status"] == "available"
        assert len(result["scenario_rows"]) == 9
        assert result["history"][-1]["page"] == 131
        book = openpyxl.load_workbook(BytesIO(export_workbook(
            Settings(company="Excel only", latest_year=2025), values, [], scenarios=result)))
        assert book["Scenarios"]["D7"].value == "=$D$4*(1+C7)^(B7-$B$4)"
        assert book["Scenarios"]["D4"].value == 50503
    else:
        assert [(c["value"], c["page"]) for c in selected("balance_sheet_8")] == [("1617", 2)]
        assert [(c["value"], c["page"]) for c in selected("balance_sheet_17")] == [("11977", 2)]
        added, _ = enrich_financial_evidence(pages, Settings(company="Excel only", latest_year=2025))
        assert [(c["value"], c["page"]) for c in added if c["metric_id"] == "income_statement_22"
                and c["year"] == 2025] == [("227", 7)]
        cover = kpi_coverage(values + added, Settings(company="Excel only", latest_year=2025))
        assert next(c for c in cover if c["id"] == "interest_coverage" and c["year"] == 2025)["status"] == "available"
        assert operating_scenarios(values, Settings(company="Excel only", latest_year=2025))["status"] == "insufficient_history"


def test_gap_search_exposes_leads_without_guessing_missing_amounts(tmp_path, settings):
    store = Store(tmp_path)
    pages = [{"sha256": "b" * 64, "file": "filing.pdf", "page": page,
              "text": f"2025 Liquid investments {page * 10}" if page == 7 else "2025 unrelated text",
              "panels": [{"side": "full", "text": f"2025 Liquid investments {page * 10}"
                          if page == 7 else "2025 unrelated text", "rows": [], "note_rows": [],
                          "headers": [{"year": 2025}] if page == 7 else []}]}
             for page in range(1, 10)]
    store.index_pages("case", pages)
    coverage = [{"missing": [{"metric_id": "balance_sheet_9", "year": 2025}], "ambiguous": []}]
    results = investigate_gaps(pages, coverage, settings,
                               lambda term, **kwargs: store.search("case", term, **kwargs))
    entry = results[0]
    assert entry["searched_pages"] == 9
    assert entry["state"] == "related_evidence_needs_review"
    assert entry["leads"][0]["page"] == 7
    assert not entry.get("candidate")


def test_forecast_requires_history_or_explicit_assumption(settings):
    observations = [{"metric_id": "income_statement_8", "year": 2025,
                     "value": "100", "file": "annual.pdf", "page": 5}]
    assert operating_scenarios(observations, settings)["status"] == "insufficient_history"
    projection = operating_scenarios(observations, settings, shock_pp=10,
                                      baseline_override=Decimal("0.04"))
    assert projection["status"] == "available"
    assert projection["scenario_rows"][4]["revenue"] == "108.1600"
    assert projection["scenario_rows"][4]["ebit"] is None
    with pytest.raises(ValueError):
        operating_scenarios(observations, settings, shock_pp=21)


def test_adidas_financial_notes_remain_distinct_from_primary_statements():
    path_text = os.environ.get("WORKBENCH_ADIDAS_NOTES_PDF")
    if not path_text:
        pytest.skip("Set WORKBENCH_ADIDAS_NOTES_PDF to the published notes")
    path = Path(path_text)
    pages, _ = read_pdf(path, path.name)
    assert len(pages) > 100
    assert not statement_plan(pages)[0]
    note_rows = [row for page in pages for panel in page["panels"]
                 for row in panel.get("note_rows", [])]
    assert len(note_rows) >= 200
    assert all(row["values"] and row["currency"] == "EUR" for row in note_rows)


def test_financial_statements_at_the_end_are_not_dropped():
    pages = []
    for n in range(180):
        rows = ([{"row_id": f"{n}:full:0", "label": "Revenue", "statement": "income"}]
                if n in (120, 179) else [])
        pages.append({"file": "annual.pdf", "page": n + 1,
                      "panels": [{"side": "full", "text": "Audit narrative", "rows": rows}]})
    panels, tasks = statement_plan(pages)
    assert {(p["page"], p["side"]) for p in panels} == {(121, "full"), (180, "full")}
    assert len(tasks) == 2


def test_mapping_rejects_combined_balances_and_cannot_change_numbers(settings, page):
    header = "31/12/2025"
    rows = []
    for i, (label, raw) in enumerate([
        ("Cash and cash equivalents", "2,500"),
        ("Trade and other receivables", "6,800"),
    ]):
        rows.append({"row_id": f"1:full:{i}", "page": 1, "file": "annual.pdf",
                     "panel": "full", "statement": "balance", "section": "Current assets",
                     "label": label, "quote": f"{label} {raw}", "scale": 1000, "currency": "EUR",
                     "values": [{"year": 2025, "scope": "consolidated", "raw_value": raw,
                                 "header": header, "x": 260}]})
    page = {**page, "text": page["text"] + " " + header + " " +
            " ".join(row["quote"] for row in rows)}
    requests = []

    async def fake(messages, structured=False, **kwargs):
        requests.append(messages)
        assert kwargs["schema"].__name__ == "RowMappings"
        assert "2,500" not in json.dumps(messages) and "6,800" not in json.dumps(messages)
        return json.dumps({"mappings": [{"row_id": rows[1]["row_id"],
                                          "metric_id": "balance_sheet_10"}]})

    candidates, rejected = asyncio.run(extract(
        [page], settings, lambda a, b: None, complete=fake,
        plan=([{"rows": rows}], [rows]),
    ))
    assert len(requests) == 1
    assert [(c["metric_id"], c["value"]) for c in candidates] == [("balance_sheet_8", "2.500")]
    assert len(rejected) == 1 and "Combined" in rejected[0]["reason"]


def test_row_mapping_calls_are_bounded():
    rows = [{"row_id": str(i), "label": f"Other financial item {i}", "statement": "balance"}
            for i in range(36)]
    plan = ([{"rows": rows}], [rows[:16], rows[16:32], rows[32:]])
    tasks = mapping_tasks(plan)
    assert [len(task) for task in tasks] == [16, 16, 4]


@pytest.mark.parametrize("lang", ["en", "el"])
def test_export_preserves_template_and_blanks(settings, page, lang):
    settings.language = lang
    c = validate_fact(fact(), page, settings)
    d = {**c, "manual": False}
    wb = openpyxl.load_workbook(BytesIO(export_workbook(settings, [d], [])))
    template = openpyxl.load_workbook(
        Path("financial_workbench/templates") / f"{lang}.xlsx"
    )
    assert wb.sheetnames[:24] == template.sheetnames
    s = METRICS["income_statement_8"]["sheet"][lang]
    assert wb[s]["I8"].value == 1.2345 and wb[s]["H8"].value is None
    assert "annual.pdf" in wb[s]["L8"].value
    for original in template:
        actual = wb[original.title]
        assert len(actual._charts) == len(original._charts)
        for row in original:
            for cell in row:
                if cell.data_type == "f":
                    assert actual[cell.coordinate].value == cell.value
    assert wb.calculation.fullCalcOnLoad
    assert len(wb.worksheets[-1]["A"]) == 1 + len(CATALOG) * 6
    draft = openpyxl.load_workbook(BytesIO(export_workbook(settings, [d], [], reviewed=False)))
    expected = "Unreviewed" if lang == "en" else "Μη ελεγμένο"
    assert expected in [r[3].value for r in draft.worksheets[-1].iter_rows(min_row=2)]
    assert ("DRAFT" if lang == "en" else "ΠΡΟΣΧΕΔΙΟ") in draft.worksheets[0]["K7"].value


def test_formula_injection_and_zero(settings):
    settings.company = '=HYPERLINK("evil")'
    decisions = [
        {
            "metric_id": "income_statement_8",
            "year": 2025,
            "value": "0",
            "manual": True,
            "note": "=1+1",
        }
    ]
    wb = openpyxl.load_workbook(BytesIO(export_workbook(settings, decisions, [])))
    assert wb["Setup"]["E6"].data_type == "s"
    assert wb["Income Statement"]["I8"].value == 0
    assert not any(c.data_type == "f" for row in wb["Export Review"] for c in row)


def test_reconciliation_missing_and_mismatch(settings):
    data = [
        {"metric_id": k, "year": 2025, "value": v}
        for k, v in [
            ("balance_sheet_29", "100"),
            ("balance_sheet_45", "60"),
            ("balance_sheet_52", "39"),
        ]
    ]
    checks = reconcile(data, settings)
    c = next(
        c
        for c in checks
        if c["year"] == 2025 and c["check"] == "Assets = liabilities + equity"
    )
    assert c["status"] == "mismatch" and c["difference"] == "1"
    assert any(c["status"] == "missing" for c in checks)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("GROQ_API_KEY", "test-placeholder")
    monkeypatch.delenv("WORKBENCH_API_KEY", raising=False)
    with TestClient(create_app(tmp_path, start_worker=False)) as c:
        yield c


def upload(client, settings):
    d = pymupdf.open()
    d.new_page().insert_text((40, 60), "Annual report 2025 narrative and audit background.")
    p = d.new_page(width=650, height=850)
    p.insert_text((40, 40), "Statement of Profit or Loss")
    p.insert_text((40, 60), "In 000's Euros")
    p.insert_text((355, 75), "GROUP", fontsize=8)
    p.insert_text((505, 75), "COMPANY", fontsize=8)
    for x, t in ((330, "31/12/2025"), (405, "31/12/2024"),
                 (480, "31/12/2025"), (555, "31/12/2024")):
        p.insert_text((x, 90), t, fontsize=8)
    p.insert_text((40, 130), "Revenue")
    for x, t in ((330, "100"), (405, "90"), (480, "80"), (555, "70")):
        p.insert_text((x, 130), t, fontsize=8)
    pdf = d.tobytes()
    d.close()
    response = client.post(
        "/api/jobs",
        data={"settings": settings.model_dump_json()},
        files=[("files", ("report.pdf", pdf, "application/pdf"))],
    )
    assert response.status_code == 202
    return response.json()


def test_api_upload_review_export_persistence(client, settings, page):
    job = upload(client, settings)
    jid = job["id"]
    assert client.get(f"/api/jobs/{jid}/workbook").status_code == 409
    stored = client.app.state.store.get(jid)
    c = validate_fact(fact(), page, settings)
    stored.update(status="review", candidates=[c])
    client.app.state.store.put(stored)
    payload = {
        "decisions": [
            {"metric_id": c["metric_id"], "year": 2025, "candidate_id": c["id"]}
        ]
    }
    assert client.put(f"/api/jobs/{jid}/review", json=payload).status_code == 200
    assert client.get(f"/api/jobs/{jid}").json()["decisions"][0]["value"] == "1.2345"
    for lang in ["en", "el"]:
        response = client.get(f"/api/jobs/{jid}/workbook?language={lang}")
        assert response.status_code == 200
        wb = openpyxl.load_workbook(BytesIO(response.content))
        assert wb[METRICS[c["metric_id"]]["sheet"][lang]]["I8"].value == 1.2345
    assert client.get(f"/api/jobs/{jid}/workbook?language=xx").status_code == 422
    assert client.get(f"/api/jobs/{jid}/audit").status_code == 200
    assert client.delete(f"/api/jobs/{jid}").status_code == 204
    assert client.get(f"/api/jobs/{jid}").status_code == 404


def test_invalid_upload(client, settings):
    response = client.post(
        "/api/jobs",
        data={"settings": settings.model_dump_json()},
        files=[("files", ("a.pdf", b"bad", "application/pdf"))],
    )
    assert response.status_code == 422
    assert client.get("/api/jobs").json() == []


def test_auth(client, monkeypatch):
    monkeypatch.setenv("WORKBENCH_API_KEY", "private")
    assert client.get("/api/jobs").status_code == 401
    assert (
        client.get("/api/catalog", headers={"X-API-Key": "private"}).status_code == 200
    )


def test_invalid_review(client, settings):
    job = upload(client, settings)
    jid = job["id"]
    stored = client.app.state.store.get(jid)
    stored["status"] = "review"
    client.app.state.store.put(stored)
    for d in [
        {"metric_id": "income_statement_8", "year": 2025, "candidate_id": "fake"},
        {"metric_id": "income_statement_8", "year": 2025, "manual_value": "10"},
        {
            "metric_id": "market_inputs_18",
            "year": 2025,
            "manual_value": "25",
            "note": "WACC",
        },
        {
            "metric_id": "income_statement_8",
            "year": 1900,
            "manual_value": "10",
            "note": "report",
        },
        {
            "metric_id": "income_statement_8",
            "year": 2025,
            "manual_value": "NaN",
            "note": "report",
        },
    ]:
        assert (
            client.put(f"/api/jobs/{jid}/review", json={"decisions": [d]}).status_code
            == 422
        )
    d = {
        "metric_id": "market_inputs_18",
        "year": 2025,
        "manual_value": "0.08",
        "note": "Analyst scenario, not reported",
    }
    assert (
        client.put(f"/api/jobs/{jid}/review", json={"decisions": [d]}).status_code
        == 200
    )
    assert (
        client.put(f"/api/jobs/{jid}/review", json={"decisions": [d, d]}).status_code
        == 422
    )


def test_worker_pdf_to_ready_without_network(tmp_path, monkeypatch, settings):
    import time

    import financial_workbench.api as api_module

    monkeypatch.setenv("GROQ_API_KEY", "test-placeholder")
    monkeypatch.delenv("WORKBENCH_API_KEY", raising=False)

    async def fake_complete(*args, **kwargs):
        raise AssertionError("Exact printed labels must not require a Groq request")

    async def extraction(pages, config, progress, plan=None):
        return await extract(pages, config, progress, complete=fake_complete, plan=plan)

    monkeypatch.setattr(api_module, "extract", extraction)
    with TestClient(create_app(tmp_path)) as client:
        job = upload(client, settings)
        for _ in range(200):
            result = client.get(f"/api/jobs/{job['id']}").json()
            if result["status"] in ("review", "failed"):
                break
            time.sleep(0.05)
        assert result["status"] == "review", result.get("error")
        assert result["candidates"][0]["value"] == "0.100"
        c = result["candidates"][0]
        assert c["column_header"] == "31/12/2025"
        evidence = client.get(f"/api/jobs/{job['id']}/evidence").json()["items"]
        assert len(evidence) == 1 and evidence[0]["row_count"] == 1
        rows = client.get(f"/api/jobs/{job['id']}/evidence?page=2").json()["items"][0]["rows"]
        assert rows[0]["values"][0]["raw_value"] == "100"
        assert any(hit["page"] == 2 for hit in client.get(
            f"/api/jobs/{job['id']}/search", params={"q": "revenue"}).json()["items"])
        assert client.get(f"/api/jobs/{job['id']}/audit").json()["statement_evidence"][0]["rows"]
        assert client.get(f"/api/jobs/{job['id']}/workbook").status_code == 409
        draft_response = client.get(f"/api/jobs/{job['id']}/workbook?language=en&draft=true")
        assert draft_response.status_code == 200
        assert f"financial-draft-en-2025-{job['id'][:8]}.xlsx" in draft_response.headers["content-disposition"]
        draft_wb = openpyxl.load_workbook(BytesIO(draft_response.content))
        assert draft_wb["Income Statement"]["I8"].value == .1
        assert draft_wb["Income Statement"]["H8"].value == .09
        statuses = [r[3].value for r in draft_wb["Export Review"].iter_rows(min_row=2)]
        assert statuses.count("Unreviewed") == 2
        assert "DRAFT" in draft_wb.worksheets[0]["K7"].value
        assert client.put(f"/api/jobs/{job['id']}/review", json={"decisions": []}).status_code == 422
        assert client.get(
            f"/api/jobs/{job['id']}/documents/{c['document_id']}"
        ).content.startswith(b"%PDF-")
        assert (
            client.put(
                f"/api/jobs/{job['id']}/review",
                json={
                    "decisions": [
                        {
                            "metric_id": c["metric_id"],
                            "year": 2025,
                            "candidate_id": c["id"],
                        }
                    ]
                },
            ).status_code
            == 200
        )
        wb = openpyxl.load_workbook(
            BytesIO(client.get(f"/api/jobs/{job['id']}/workbook").content)
        )
        assert wb["Income Statement"]["I8"].value == 0.1
        assert "DRAFT" not in wb.worksheets[0]["K7"].value


def test_question_isolation(client, settings, page, monkeypatch):
    import financial_workbench.api as api_module

    one = upload(client, settings)
    two = upload(client, settings)
    for job, body in [
        (one, page),
        (two, {**page, "text": "Other company confidential SECRET_987 revenue 9900"}),
    ]:
        stored = client.app.state.store.get(job["id"])
        stored.update(status="review", pages=[body])
        client.app.state.store.put(stored)
        client.app.state.store.index_pages(job["id"], [body])

    async def fake_complete(messages, structured=False, **kwargs):
        assert "SECRET_987" not in json.dumps(messages)
        assert "Greek" in messages[0]["content"]
        return "Έσοδα με αναφορά [1]."

    monkeypatch.setattr(api_module, "completion", fake_complete)
    response = client.post(
        f"/api/jobs/{one['id']}/ask",
        json={"question": "revenue 2025", "language": "el"},
    )
    assert response.status_code == 200
    assert response.json()["sources"][0]["file"] == "annual.pdf"


def test_search_index_is_persistent_and_job_scoped(tmp_path):
    a, b = "a" * 32, "b" * 32
    store = Store(tmp_path)
    store.index_pages(a, [{"sha256": "1" * 64, "file": "first.pdf", "page": 179,
                           "text": "Revenue 100 on the final page",
                           "panels": [{"side": "right", "text": "Revenue 100 on the final page",
                                       "rows": [{"label": "Revenue", "quote": "Revenue 100"}]}]}])
    store.index_pages(b, [{"sha256": "2" * 64, "file": "other.pdf", "page": 1,
                           "text": "Revenue SECRET_987"}])
    assert len(Store(tmp_path).search(a, "revenue")) == 2
    assert all("SECRET_987" not in x["text"] for x in store.search(a, "revenue"))
    assert store.search(a, 'revenue" OR SECRET_987')
    store.delete(a)
    assert not Store(tmp_path).search(a, "revenue")
    assert Store(tmp_path).search(b, "revenue")


def test_search_returns_match_from_late_in_a_long_pdf_page(tmp_path):
    store = Store(tmp_path)
    store.index_pages("x", [{"sha256": "f" * 64, "file": "large.pdf", "page": 190,
                             "text": "narrative " * 2000 + "Final debt maturities are 2030."}])
    matches = store.search("x", "debt maturities")
    assert len(matches) == 1 and "Final debt maturities" in matches[0]["text"]
    assert len(matches[0]["text"]) < 1200


@pytest.mark.skipif(not os.getenv("WORKBENCH_REPORT_PDF"), reason="Requires locally supplied reference PDF")
def test_real_annual_report_statement_columns_and_grounding(settings):
    pdf = Path(os.environ["WORKBENCH_REPORT_PDF"])
    pages, warnings = read_pdf(pdf, pdf.name)
    panels, batches = statement_plan(pages)
    assert len(pages) == 187
    assert {(p["page"], p["side"]) for p in panels} == {
        (125, "right"), (126, "left"), (126, "right"),
        (127, "left"), (128, "right"), (129, "left"),
    }
    assert len(batches) <= 16
    assert any("Revenue" in r["label"] for p in panels for r in p["rows"])
    assert not any(p["page"] == 183 for p in panels)

    async def empty_map(messages, structured=False, **kwargs):
        assert kwargs["schema"].__name__ == "RowMappings"
        assert not any(raw in json.dumps(messages) for raw in ("11,482,478", "8,039,778"))
        return '{"mappings":[]}'

    candidates, rejected = asyncio.run(extract(
        pages, settings, lambda *_: None, complete=empty_map, plan=(panels, batches)))
    figures = {(c["metric_id"], c["year"]): c for c in candidates}
    for metric, year, amount, source_page, side in (
        ("income_statement_8", 2025, "11482.478", 125, "right"),
        ("income_statement_26", 2025, "176.435", 125, "right"),
        ("balance_sheet_29", 2025, "8039.778", 126, "right"),
        ("cash_flow_statement_18", 2025, "799.916", 128, "right"),
        ("cash_flow_statement_27", 2025, "-449.332", 129, "left"),
    ):
        candidate = figures[(metric, year)]
        assert (candidate["value"], candidate["page"], candidate["panel"]) == (
            amount, source_page, side)
        assert candidate["mapping_source"] == "exact label"
        assert candidate["raw_value"] in candidate["quote"]
    assert all(c["scope"] == "consolidated" for c in candidates)
    assert all("trade and other" not in c["row_label"].lower()
               for c in candidates if c["metric_id"] in ("balance_sheet_10", "balance_sheet_31"))


@pytest.mark.skipif(not os.getenv("WORKBENCH_EVIDENCE_JSON"), reason="Requires a local evidence export")
def test_user_evidence_can_produce_nonempty_draft():
    source = json.loads(Path(os.environ["WORKBENCH_EVIDENCE_JSON"]).read_text(encoding="utf-8"))
    assert source["status"] == "review" and source["decisions"] == []
    selected = provisional_decisions(source["candidates"])
    assert len(selected) > 100
    settings = Settings(**source["settings"])
    wb = openpyxl.load_workbook(BytesIO(export_workbook(settings, selected, [], reviewed=False)))
    assert wb["Income Statement"]["I8"].value == 11482.478
    assert wb["Balance Sheet"]["I29"].value == 8039.778
    assert sum(row[3].value == "Unreviewed" for row in wb["Export Review"].iter_rows(min_row=2)) == len(selected)


@pytest.mark.skipif(not os.getenv("WORKBENCH_REPORT_PDF") or not os.getenv("WORKBENCH_EVIDENCE_JSON"),
                    reason="Requires the locally supplied annual report and evidence export")
def test_supplied_json_recovers_note_values_and_preserves_distinct_kpi_definitions(tmp_path):
    pdf = Path(os.environ["WORKBENCH_REPORT_PDF"])
    supplied = json.loads(Path(os.environ["WORKBENCH_EVIDENCE_JSON"]).read_text(encoding="utf-8"))
    pages, _ = read_pdf(pdf, pdf.name)
    settings = Settings(**supplied["settings"])
    store = Store(tmp_path)
    store.index_pages(supplied["id"], pages)
    additions, reported = enrich_financial_evidence(
        pages, settings, search=lambda q, **kw: store.search(supplied["id"], q, **kw))
    found = {(c["metric_id"], c["year"]): c for c in additions}
    for metric, value, page in (
        ("balance_sheet_10", "662.237", 155),
        ("income_statement_17", "286.142", 128),
        ("income_statement_22", "111.290", 145),
        ("income_statement_18", "1059.348", 7),
        ("market_inputs_7", "31.40", 9),
        ("market_inputs_9", "108.315628", 146),
    ):
        assert (found[(metric, 2025)]["value"], found[(metric, 2025)]["page"]) == (value, page)
    assert found[("income_statement_22", 2025)]["reconciliation"]["reported"] == "253.651"
    assert found[("balance_sheet_10", 2025)]["reconciliation"]["reported"] == "662.237"
    merged, removed = merge_evidence(supplied["candidates"], additions)
    assert len(removed) >= 20
    assert all(not (c["metric_id"] == "cash_flow_statement_15" and c["row_label"] == "Finance cost")
               for c in merged)
    checked = check_reported_measures(merged, reported)
    assert len(checked) == 1 and checked[0]["reconciliation"]["calculated"] == "1579.448"
    coverage = {c["id"]: c for c in kpi_coverage(merged, settings, checked) if c["year"] == 2025}
    assert coverage["quick_ex_inventory"]["status"] == "available"
    assert coverage["interest_coverage"]["value"] == str(Decimal("771.283") / Decimal("111.290"))
    assert coverage["reported_net_debt_ebitda"]["value"] == str(Decimal("1579.448") / Decimal("1059.348"))
    assert coverage["pe"]["value"] == str(Decimal("31.40") / Decimal("5.98"))
    assert coverage["net_debt"]["status"] == "missing"
    assert "balance_sheet_9" in [c["metric_id"] for c in coverage["net_debt"]["missing"]]
    assert "market_inputs_17" in [c["metric_id"] for c in coverage["roic"]["missing"]]
    selected = provisional_decisions(merged)
    selected_measures = check_reported_measures(selected, [dict(m) for m in reported])
    wb = openpyxl.load_workbook(BytesIO(export_workbook(settings, selected, [], reviewed=False,
        coverage=kpi_coverage(selected, settings, selected_measures))))
    assert wb["Income Statement"]["I17"].value == 286.142
    assert wb["Income Statement"]["I22"].value == 111.29
    assert wb["Balance Sheet"]["I10"].value == 662.237
    assert wb["Capital Calculations"]["I9"].data_type == "f"
    evidence_sheet = wb["Evidence KPIs"]
    ratio_row = next(r for r in evidence_sheet.iter_rows(min_row=2)
                     if r[0].value == 2025 and r[1].value == "Reported net debt / reported EBITDA (issuer definition)")
    assert ratio_row[2].value == f"=H{ratio_row[0].row}/I{ratio_row[0].row}"
    assert (ratio_row[7].value, ratio_row[8].value) == (1579.448, 1059.348)
    assert "p.42" in ratio_row[7].comment.text


@pytest.mark.skipif(not os.getenv("WORKBENCH_REPORT_PDF") or not os.getenv("WORKBENCH_EVIDENCE_JSON"),
                    reason="Requires the locally supplied annual report and evidence export")
def test_recheck_saved_job_without_provider_or_reupload(tmp_path, monkeypatch):
    pdf = Path(os.environ["WORKBENCH_REPORT_PDF"])
    saved = json.loads(Path(os.environ["WORKBENCH_EVIDENCE_JSON"]).read_text(encoding="utf-8"))
    saved.pop("statement_evidence", None)
    saved["files"] = [{"name": pdf.name, "path": str(pdf)}]
    saved["pages"] = []
    previously_selected = saved["candidates"][0]
    saved["decisions"] = [{**previously_selected, "manual": False, "note": "Previously reviewed"}]
    saved["status"] = "ready"
    import financial_workbench.api as workbench_api

    async def no_provider(*_args, **_kwargs):
        raise AssertionError("Refresh must not contact Groq")

    monkeypatch.setattr(workbench_api, "completion", no_provider)
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        client.app.state.store.put(saved)
        endpoint = f"/api/jobs/{saved['id']}/refresh-evidence"
        first = client.post(endpoint)
        assert first.status_code == 200, first.text
        result = first.json()
        assert result["status"] == "ready"
        assert result["decisions"][0]["id"] == previously_selected["id"]
        assert any(c["metric_id"] == "income_statement_22" and c["value"] == "111.290"
                   for c in result["candidates"])
        assert result["coverage"] and result["reported_measures"]
        inspected = client.get(f"/api/jobs/{saved['id']}/evidence", params={"page": 145}).json()["items"]
        assert any("interest on borrowings" in r["label"].lower()
                   for panel in inspected for r in panel["note_rows"])
        exported = client.get(f"/api/jobs/{saved['id']}/audit").json()
        assert any(panel["page"] == 145 for panel in exported["note_evidence"])
        again = client.post(endpoint)
        assert again.status_code == 200, again.text
        assert len(again.json()["candidates"]) == len(result["candidates"])


def test_restart_marks_inflight_failed(tmp_path, monkeypatch, settings):
    monkeypatch.setenv("GROQ_API_KEY", "test-placeholder")
    monkeypatch.delenv("WORKBENCH_API_KEY", raising=False)
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        j = upload(client, settings)
        job = client.app.state.store.get(j["id"])
        job["status"] = "extracting"
        client.app.state.store.put(job)
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        assert client.get(f"/api/jobs/{j['id']}").json()["status"] == "failed"
        assert client.post(f"/api/jobs/{j['id']}/retry").json()["status"] == "queued"


def test_restart_backfills_search_for_saved_jobs(tmp_path, monkeypatch, settings):
    monkeypatch.setenv("GROQ_API_KEY", "test-placeholder")
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        job = upload(client, settings)
        saved = client.app.state.store.get(job["id"])
        saved.update(status="failed", pages=[{"sha256": "c" * 64,
                                             "file": "annual.pdf", "page": 151,
                                             "text": "The late-page debt schedule is searchable."}])
        client.app.state.store.put(saved)
        assert not client.app.state.store.has_index(job["id"])
    with TestClient(create_app(tmp_path, start_worker=False)) as client:
        response = client.get(f"/api/jobs/{job['id']}/search", params={"q": "debt schedule"})
        assert response.status_code == 200
        assert response.json()["items"][0]["page"] == 151


def test_provider_schema_and_incomplete_output(monkeypatch):
    import httpx

    from financial_workbench.llm import ProviderError, completion

    monkeypatch.setenv("GROQ_API_KEY", "test-placeholder")
    original = httpx.AsyncClient

    def handler(request):
        body = json.loads(request.content)
        assert body["response_format"]["json_schema"]["strict"] is True
        return httpx.Response(
            200,
            json={
                "choices": [
                    {"finish_reason": "length", "message": {"content": '{"facts":'}}
                ]
            },
        )

    monkeypatch.setattr(
        httpx,
        "AsyncClient",
        lambda **kwargs: original(transport=httpx.MockTransport(handler), **kwargs),
    )
    with pytest.raises(ProviderError, match="incomplete"):
        asyncio.run(completion([{"role": "user", "content": "test"}], structured=True))


def test_provider_retry_delay_parses_daily_limit_window():
    from financial_workbench.llm import retry_delay

    class FakeResponse:
        headers = {}

        @staticmethod
        def json():
            return {
                "error": {
                    "message": "Tokens per day limit reached. Please try again in 3m42.047999999s."
                }
            }

    assert retry_delay(FakeResponse(), 0) == pytest.approx(223.048)


def test_dotted_dates_numbered_titles_and_issuer_units():
    from financial_workbench.documents import _statement_title, _unit, _year
    assert _year("31.12.2025") == 2025
    assert _statement_title("Annual Financial Report\n1.  Statement of Financial Position") == "balance"
    assert _statement_title("Annual Financial Report\n2.  Statement of Comprehensive Income") == "income"
    assert _unit("Annual Financial Report\n(AmountsinEuro) GROUP COMPANY") == ("EUR", 1)


def test_aktor_2025_primary_statements_and_note_bridges():
    """Optional real-PDF regression: run with WORKBENCH_AKTOR_PDF=/path/to/report.pdf."""
    pdf = os.environ.get("WORKBENCH_AKTOR_PDF")
    if not pdf:
        pytest.skip("Set WORKBENCH_AKTOR_PDF to AKTOR's 306-page FY2025 annual report")
    pages, _ = read_pdf(Path(pdf), Path(pdf).name)
    assert len(pages) == 306
    panels, batches = statement_plan(pages)
    assert panels and batches
    balance = [p for p in panels if p["page"] == 208 and p["statement"] == "balance"]
    income = [p for p in panels if p["page"] == 209 and p["statement"] == "income"]
    assert balance and income
    notes = [r for p in pages for x in p["panels"] for r in x.get("note_rows", [])
             if p["page"] == 264]
    assert any(r["label"].strip() == "Final trade receivables" for r in notes)
    async def no_model(messages, **kwargs):
        return '{"mappings":[]}'
    settings = Settings(company="AKTOR", latest_year=2025)
    facts, _ = asyncio.run(extract(pages, settings, lambda *_: None,
                                   complete=no_model, plan=(panels, batches)))
    for metric, amount in (("balance_sheet_8", "268.681474"),
                           ("balance_sheet_17", "1504.190595"),
                           ("balance_sheet_37", "1243.781966"),
                           ("income_statement_8", "1394.998664"),
                           ("income_statement_16", "72.650399")):
        assert amount in [c["value"] for c in facts
                          if c["metric_id"] == metric and c["year"] == 2025]
    added, _ = enrich_financial_evidence(pages, settings)
    assert any(c["metric_id"] == "balance_sheet_10" and
               c["year"] == 2025 and c["value"] == "254.775833" for c in added)
    assert any(c["metric_id"] == "income_statement_22" and
               c["year"] == 2025 and c["value"] == "39.035163" for c in added)
    book = export_workbook(settings, provisional_decisions(facts + added), [],
                           reviewed=False)
    assert book.startswith(b"PK")


@pytest.mark.parametrize("text, scale", [
    # AKTOR FY2025 p.209: "Amounts in Euro" header, round company amount 349.523.000.
    ("2. Statement of Comprehensive Income\n(AmountsinEuro) GROUP COMPANY\n"
     "Sales 6.31 1.394.998.664 1.254.923.600 349.523.000 481.696.745", 1),
    ("Statement of profit or loss\nAmounts in euros\n2025 2024\nRevenue 1,000 2,500,000", 1),
    ("Statement of financial position\nAmounts in Euro\nThe Group employs several thousand people", 1),
    ("Statement of Profit or Loss\nIn 000's Euros", 1000),
    ("Balance sheet\n€'000 2025 2024\nCash 1,000 900", 1000),
    ("Balance sheet (Amounts in EUR thousands) 2025 2024", 1000),
    ("Example SA. Consolidated. Annual 2025. EUR thousands.", 1000),
    ("Κατάσταση Χρηματοοικονομικής Θέσης (Ποσά σε χιλιάδες €)", 1000),
    ("Ισολογισμός σε χιλ. ευρώ 2025 2024", 1000),
    ("Κατάσταση Συνολικού Εισοδήματος (Ποσά σε ευρώ) 1.000.000", 1),
    ("Consolidated income statement\n€ million 2025 2024", 1000000),
    ("Consolidated income statement\n€ in millions 2025 2024", 1000000),
    ("Income statement EURm 2025 2024", 1000000),
])
def test_unit_comes_from_declaration_not_round_amounts(text, scale):
    from financial_workbench.documents import _unit
    assert _unit(text) == ("EUR", scale)


def test_undeclared_unit_is_not_guessed_from_amounts():
    from financial_workbench.documents import _unit
    assert _unit("Group EUR 1.000.000 2025 2024") is None
    assert _unit("Statement of cash flows EUR 2025 2024 Cash 1.000") is None


def _dated_page(doc, lines, rows, dates=("31.12.2025", "31.12.2024", "31.12.2025", "31.12.2024")):
    page = doc.new_page(width=650, height=850)
    for y, text in lines:
        page.insert_text((40, y), text)
    page.insert_text((355, 75), "GROUP", fontsize=8)
    page.insert_text((505, 75), "COMPANY", fontsize=8)
    for x, text in zip((330, 405, 480, 555), dates):
        page.insert_text((x, 90), text, fontsize=8)
    y = 130
    for label, note, values in rows:
        page.insert_text((40, y), label, fontsize=8)
        if note:
            # Base-14 fonts cannot draw Greek note suffixes such as 7.24α.
            page.insert_text((270, y), note, fontsize=8,
                             fontname="helv" if note.isascii() else "china-s")
        for x, text in zip((330, 405, 480, 555), values):
            page.insert_text((x, y), text, fontsize=8)
        y += 18
    return page


def _aktor_like_pdf(path, income_unit="(AmountsinEuro)", balance_rows=()):
    doc = pymupdf.open()
    _dated_page(doc, [(40, "1. Statement of Financial Position"), (60, "(AmountsinEuro)")], [
        ("Cash and cash equivalents", "7.17", ("268.681.474", "106.554.606", "121.605.950", "33.123.855")),
        ("Total Current Assets", "", ("1.504.190.595", "1.110.593.347", "229.283.911", "408.099.512")),
        *balance_rows,
    ])
    _dated_page(doc, [(40, "2. Statement of Comprehensive Income"), (60, income_unit)], [
        ("Sales", "6.31", ("1.394.998.664", "1.254.923.600", "349.523.000", "481.696.745")),
        ("Operating results", "", ("72.650.399", "61.143.368", "(45.993.376)", "(333.600)")),
    ])
    doc.new_page().insert_text((40, 60), "5. General information")
    _dated_page(doc, [(40, "7.9 Trade and other receivables")], [
        ("Trade receivables", "", ("297.574.264", "359.358.867", "1.080", "124.697.047")),
    ])
    doc.save(path)
    doc.close()


def test_round_company_amount_keeps_euro_statement_in_euros(tmp_path, settings):
    """AKTOR FY2025 draft regression: p.209 figures were exported 1,000x too large."""
    pdf = tmp_path / "aktor-like.pdf"
    _aktor_like_pdf(pdf)
    pages, warnings = read_pdf(pdf, pdf.name)
    assert not [w for w in warnings if "different units" in w]
    income = [r for x in pages[1]["panels"] for r in x["rows"]]
    assert {r["scale"] for r in income} == {1}
    assert income[0]["unit_source"] == "declared on page 2"
    note = [r for x in pages[3]["panels"] for r in x.get("note_rows", [])]
    assert note and note[0]["scale"] == 1
    assert note[0]["unit_source"] == "inherited from page 2"

    async def no_model(messages, **kwargs):
        raise AssertionError("Exact labels must not need the provider")

    facts, _ = asyncio.run(extract(pages, settings, lambda *_: None, complete=no_model))
    values = {(c["metric_id"], c["year"]): Decimal(c["value"]) for c in facts}
    assert values[("income_statement_8", 2025)] == Decimal("1394.998664")
    assert values[("income_statement_8", 2024)] == Decimal("1254.9236")
    assert values[("income_statement_16", 2025)] == Decimal("72.650399")
    assert values[("balance_sheet_8", 2025)] == Decimal("268.681474")


def test_statements_with_different_declared_units_are_flagged(tmp_path):
    pdf = tmp_path / "mixed.pdf"
    _aktor_like_pdf(pdf, income_unit="(Amounts in EUR thousands)")
    _, warnings = read_pdf(pdf, pdf.name)
    assert any("different units" in w and "x1,000 on page(s) 2" in w for w in warnings)


@pytest.mark.parametrize("text, kind", [
    ("Annual Financial Report\n4.  Cash Flow Statement", "cash"),
    ("4. Statement of Cash Flows (indirect method)", "cash"),
    ("3. Statement of Changes in Equity", "equity"),
    ("Consolidated statement of changes in shareholders' equity", "equity"),
    ("ΚΑΤΑΣΤΑΣΗ ΤΑΜΕΙΑΚΩΝ ΡΟΩΝ", "cash"),
    ("Κατάσταση Χρηματοοικονομικής Θέσης", "balance"),
    ("Κατάσταση Συνολικού Εισοδήματος", "income"),
    ("Κατάσταση Μεταβολών Ιδίων Κεφαλαίων", "equity"),
])
def test_numbered_and_greek_statement_titles(text, kind):
    from financial_workbench.documents import _statement_title
    assert _statement_title(text) == kind


@pytest.mark.parametrize("header, year", [
    ("31/12/2025", 2025), ("31.12.2025", 2025),
    ("1/1-31/12/2025", 2025), ("01.01-31.12.2025", 2025),
    ("1.1.-31.12.2025", 2025), ("01.01.2025-31.12.2025", 2025),
    ("01.07.2024-30.06.2025", 2025),
])
def test_year_end_and_period_column_headers(header, year):
    from financial_workbench.documents import DATE_PATTERN, _year
    assert DATE_PATTERN.fullmatch(header) and _year(header) == year


def test_model_cannot_shift_labels_onto_neighbouring_metrics(tmp_path, settings):
    """AKTOR FY2025 draft: Groq shifted three liability rows by one metric."""
    doc = pymupdf.open()
    _dated_page(doc, [(40, "1. Statement of Financial Position"), (60, "(AmountsinEuro)")], [
        ("Contractual liabilities", "7.12", ("109.036.041", "138.929.845", "-", "76.420.002")),
        ("Current income tax liabilities", "", ("5.539.829", "10.395.306", "-", "1.240.668")),
        ("Current provisions for other liabilities and expenses", "7.25",
         ("24.573.779", "-", "-", "-")),
        ("Current liabilities to third parties", "7.24α", ("-", "70.179.158", "-", "-")),
    ])
    pdf = tmp_path / "liabilities.pdf"
    doc.save(pdf)
    doc.close()
    pages, _ = read_pdf(pdf, pdf.name)
    rows = [r for x in pages[0]["panels"] for r in x["rows"]]
    shifted = dict(zip([r["row_id"] for r in rows], (
        "balance_sheet_34", "balance_sheet_35", "balance_sheet_36", "balance_sheet_31")))

    async def shifting_model(messages, **kwargs):
        ids = [r["row_id"] for r in json.loads(messages[1]["content"])["rows"]]
        return json.dumps({"mappings": [{"row_id": i, "metric_id": shifted[i]} for i in ids]})

    with pytest.raises(ValueError, match="no grounded figures"):
        asyncio.run(extract(pages, settings, lambda *_: None, complete=shifting_model))

    correct = dict(zip([r["row_id"] for r in rows], (
        "balance_sheet_36", "balance_sheet_34", "balance_sheet_35", "balance_sheet_36")))

    async def correct_model(messages, **kwargs):
        ids = [r["row_id"] for r in json.loads(messages[1]["content"])["rows"]]
        return json.dumps({"mappings": [{"row_id": i, "metric_id": correct[i]} for i in ids]})

    facts, rejected = asyncio.run(extract(pages, settings, lambda *_: None,
                                          complete=correct_model))
    values = {(c["metric_id"], c["year"]): Decimal(c["value"]) for c in facts}
    assert values[("balance_sheet_34", 2025)] == Decimal("5.539829")
    assert values[("balance_sheet_35", 2025)] == Decimal("24.573779")
    # "Current liabilities to third parties" names no payable, so it cannot be trade payables.
    assert any(r["row_label"] == "Current liabilities to third parties"
               and r["metric_id"] == "balance_sheet_31" for r in
               asyncio.run(extract(pages, settings, lambda *_: None,
                                   complete=shifting_model_31(rows)))[1])


def shifting_model_31(rows):
    async def model(messages, **kwargs):
        return json.dumps({"mappings": [
            {"row_id": rows[1]["row_id"], "metric_id": "balance_sheet_34"},
            {"row_id": rows[3]["row_id"], "metric_id": "balance_sheet_31"}]})
    return model


def test_saved_model_mappings_are_rechecked_without_provider():
    saved = {"id": "a", "metric_id": "balance_sheet_34", "year": 2025, "value": "109.036041",
             "document_id": "d", "page": 208, "panel": "full",
             "mapping_source": "Groq label mapping", "row_label": "Contractual liabilities",
             "section": "Current liabilities"}
    kept = saved | {"id": "b", "metric_id": "balance_sheet_34", "value": "5.539829",
                    "row_label": "Current income tax liabilities"}
    merged, removed = merge_evidence([saved, kept], [])
    assert [f["id"] for f in merged] == ["b"] and [f["id"] for f in removed] == ["a"]


def test_refresh_corrects_saved_scale_and_mappings_without_provider(tmp_path, monkeypatch):
    """Repair a saved AKTOR-style job in place: no new Groq request, no re-upload."""
    import hashlib
    import financial_workbench.api as workbench_api

    combined = tmp_path / "Report-2025Y-AKTOR-GROUP.pdf"
    _aktor_like_pdf(combined, balance_rows=(
        ("Committed deposit accounts", "7.17a", ("54.976.214", "42.382.864", "100.000", "3.282.005")),
        ("Contractual liabilities", "7.12", ("109.036.041", "138.929.845", "-", "76.420.002")),
    ))
    digest = hashlib.sha256(combined.read_bytes()).hexdigest()
    settings = Settings(company="aktor", latest_year=2025)
    old = lambda metric, value, label, page, source="exact label": {  # noqa: E731
        "id": f"{metric}-{page}", "metric_id": metric, "year": 2025, "value": value,
        "document_id": digest, "file": combined.name, "page": page, "panel": "full",
        "row_label": label, "section": "", "mapping_source": source, "status": "candidate"}
    wrong_revenue = old("income_statement_8", "1394998.664", "Sales", 2)
    saved = {
        "id": "f" * 32, "status": "ready", "settings": settings.model_dump(),
        "files": [{"name": combined.name, "path": str(combined)}], "pages": [],
        "candidates": [wrong_revenue,
                       old("balance_sheet_15", "54.976214", "Committed deposit accounts 7.17a", 1,
                           "Groq label mapping"),
                       old("balance_sheet_34", "109.036041", "Contractual liabilities", 1,
                           "Groq label mapping")],
        "rejected": [], "decisions": [{**wrong_revenue, "manual": False, "note": ""}],
        "checks": [], "reported_measures": [], "coverage": [], "investigations": [],
        "scenarios": None, "warnings": [], "progress": {"done": 0, "total": 0}, "error": None,
    }

    async def no_provider(*_args, **_kwargs):
        raise AssertionError("Refresh must not contact Groq")

    monkeypatch.setattr(workbench_api, "completion", no_provider)
    monkeypatch.setattr("financial_workbench.engine.completion", no_provider)
    with TestClient(create_app(tmp_path / "data", start_worker=False)) as client:
        client.app.state.store.put(saved)
        response = client.post(f"/api/jobs/{saved['id']}/refresh-evidence")
        assert response.status_code == 200, response.text
        job = response.json()
    values = {(c["metric_id"], c["year"]): Decimal(c["value"]) for c in job["candidates"]}
    assert values[("income_statement_8", 2025)] == Decimal("1394.998664")
    assert values[("balance_sheet_15", 2025)] == Decimal("54.976214")
    assert ("balance_sheet_34", 2025) not in values
    assert job["status"] == "review" and job["decisions"] == []
    assert any("Revenue 2025" in w and "1394998.664" in w for w in job["warnings"])
    assert any(r["row_label"] == "Contractual liabilities" and r["metric_id"] == "balance_sheet_34"
               for r in job["rejected"])
